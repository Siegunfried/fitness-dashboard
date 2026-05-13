"""从 data/training-plan.json 生成 ToDoList.docx 到桌面"""
import json, os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_DIR = Path(__file__).resolve().parent.parent
JSON_PATH = PROJECT_DIR / 'data' / 'training-plan.json'
DOCX_PATH = Path.home() / 'Desktop' / 'ToDoList.docx'

with open(JSON_PATH, 'r', encoding='utf-8') as f:
    data = json.load(f)

GREEN = RGBColor(0x1d, 0xb9, 0x54); GRAY = RGBColor(0x99, 0x99, 0x99)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)
style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(1); style.paragraph_format.space_before = Pt(0)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(data.get('title', 'TodoList'))
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = GREEN

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('训练计划 · 周一至周日')
r.font.size = Pt(9); r.font.color.rgb = GRAY
p.paragraph_format.space_after = Pt(10)

def add(text, indent=0, bold=False, color=None, size=None, sb=0):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(indent * 1.0)
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(sb)
    r = p.add_run(text); r.font.name = 'Calibri'
    if bold: r.bold = True
    if color: r.font.color.rgb = color
    if size: r.font.size = size

for entry in data.get('week', []):
    day = entry.get('day', ''); label = entry.get('label', '')
    title = f'{day} — {label}' if label else day
    add(title, indent=0, bold=True, color=GREEN, size=Pt(13), sb=8)
    for task in entry.get('tasks', []):
        add(task, indent=1, bold=False, size=Pt(10.5))

add('─' * 50, indent=0, color=RGBColor(0x66,0x66,0x66), size=Pt(6), sb=14)
add('数据源：training-plan.json  |  修改 JSON 后运行此脚本刷新', indent=0, color=GRAY, size=Pt(8))

doc.save(DOCX_PATH)
print(f'Done -> {DOCX_PATH}')
